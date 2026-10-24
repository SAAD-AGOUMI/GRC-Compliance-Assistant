"""
Analyse d'écarts complète : évalue toutes les mesures (ou une sélection),
calcule les taux de conformité, sauvegarde les résultats en JSON et
exporte la matrice de conformité en Excel.

Usage en ligne de commande (dans le conteneur agent) :
    python Compliance/run_audit.py                    # toutes les mesures
    python Compliance/run_audit.py --ids 8.13 5.30    # quelques mesures
    python Compliance/run_audit.py --themes "Mesures technologiques"
"""

import argparse
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from Compliance.evaluate_control import (  # noqa: E402
    CONFORME, ERREUR, NON_CONFORME, NON_TRAITEE, PARTIEL,
    evaluate_control, load_referentiel,
)

RESULTS_DIR = Path(__file__).resolve().parent / "resultats"
ORDRE_STATUTS = [CONFORME, PARTIEL, NON_CONFORME, NON_TRAITEE, ERREUR]
COULEURS = {  # fond, texte (lisibles à l'impression)
    CONFORME: ("C6EFCE", "006100"),
    PARTIEL: ("FFEB9C", "9C5700"),
    NON_CONFORME: ("FFC7CE", "9C0006"),
    NON_TRAITEE: ("D9D9D9", "3F3F3F"),
    ERREUR: ("E4DFEC", "5B2C6F"),
}


# ---------------------------------------------------------
# Exécution
# ---------------------------------------------------------
def select_mesures(ids=None, themes=None):
    mesures = load_referentiel()["mesures"]
    if ids:
        mesures = [m for m in mesures if m["id"] in set(ids)]
    if themes:
        mesures = [m for m in mesures if m["theme"] in set(themes)]
    return mesures


def run_audit(ids=None, themes=None, progress=None, **eval_kwargs):
    """
    Évalue les mesures sélectionnées.
    progress : fonction optionnelle appelée avec (index, total, mesure) avant chaque mesure.
    """
    mesures = select_mesures(ids, themes)
    results = []
    for i, m in enumerate(mesures):
        if progress:
            progress(i, len(mesures), m)
        try:
            results.append(evaluate_control(m["id"], **eval_kwargs))
        except Exception as e:  # une mesure en échec n'arrête pas l'audit
            results.append({
                "id": m["id"], "theme": m["theme"], "titre": m["titre"],
                "nis2": m.get("nis2", []), "criteres": m["criteres"],
                "statut": ERREUR, "statut_llm": None,
                "criteres_couverts": [], "criteres_manquants": m["criteres"],
                "preuves": [], "constats": [], "sources": [], "contradictions": "", "ecart": "", "recommandation": "",
                "avertissements": [f"Exception : {e}"], "sources_consultees": [], "duree_s": 0,
            })
    if progress:
        progress(len(mesures), len(mesures), None)
    return results


# ---------------------------------------------------------
# Statistiques
# ---------------------------------------------------------
def _taux(counter: Counter):
    evalues = sum(v for k, v in counter.items() if k != ERREUR)
    if evalues == 0:
        return 0.0, 0.0
    strict = 100 * counter[CONFORME] / evalues
    pondere = 100 * (counter[CONFORME] + 0.5 * counter[PARTIEL]) / evalues
    return round(strict, 1), round(pondere, 1)


def _ligne(nom, counter):
    strict, pondere = _taux(counter)
    return {
        "Périmètre": nom,
        "Mesures": sum(counter.values()),
        **{s.capitalize(): counter[s] for s in ORDRE_STATUTS},
        "Taux de conformité (%)": strict,
        "Taux pondéré (%)": pondere,
    }


def compute_stats(results):
    """
    Taux de conformité = mesures conformes / mesures évaluées.
    Taux pondéré = (conformes + 0,5 × partiellement conformes) / mesures évaluées.
    Les erreurs d'évaluation sont exclues du dénominateur.
    """
    globales = _ligne("Global", Counter(r["statut"] for r in results))

    par_theme = []
    for theme in dict.fromkeys(r["theme"] for r in results):
        par_theme.append(_ligne(theme, Counter(r["statut"] for r in results if r["theme"] == theme)))

    par_nis2 = []
    articles = sorted({a for r in results for a in r.get("nis2", [])})
    for art in articles:
        par_nis2.append(_ligne(f"NIS2 {art}", Counter(r["statut"] for r in results if art in r.get("nis2", []))))

    return {"global": globales, "par_theme": par_theme, "par_nis2": par_nis2}


# ---------------------------------------------------------
# Sauvegarde et export
# ---------------------------------------------------------
def save_results(results, stats, out_dir=RESULTS_DIR):
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"audit_{stamp}.json"
    payload = {"date": datetime.now().isoformat(timespec="seconds"),
               "referentiel": load_referentiel()["referentiel"],
               "stats": stats, "resultats": results}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def load_latest(out_dir=RESULTS_DIR):
    files = sorted(out_dir.glob("audit_*.json")) if out_dir.exists() else []
    if not files:
        return None
    return json.loads(files[-1].read_text(encoding="utf-8"))


def matrice_dataframe(results):
    rows = []
    for r in results:
        rows.append({
            "ID": r["id"],
            "Thème": r["theme"],
            "Mesure": r["titre"],
            "Statut": r["statut"],
            "Critères couverts": "\n".join(f"✓ {c}" for c in r["criteres_couverts"]),
            "Critères manquants": "\n".join(f"✗ {c}" for c in r["criteres_manquants"]),
            "Écart": r.get("ecart", ""),
            "Recommandation": r.get("recommandation", ""),
            "Sources": "\n".join(r.get("sources") or []) or "Aucune (non traitée)",
            "Preuves": "\n".join(f"{p['doc']} (p. {p['page']}) : « {p['extrait']} »" for p in r["preuves"]),
            "Constats": "\n".join(f"{c['doc']} (p. {c['page']}) : « {c['extrait']} »" for c in r.get("constats", [])),
            "Contradictions": r.get("contradictions", ""),
            "NIS2": ", ".join(r.get("nis2", [])),
            "Contrôles automatiques": "\n".join(r.get("avertissements", [])),
        })
    return pd.DataFrame(rows)


def export_excel(results, stats, path):
    """Exporte la matrice de conformité (4 onglets) au format Excel."""
    from openpyxl.styles import Alignment, Font, PatternFill

    preuves = pd.DataFrame([
        {"ID": r["id"], "Mesure": r["titre"], "Type": typ, "Critère": p["critere"],
         "Document": p["doc"], "Page": p["page"], "Extrait": p["extrait"]}
        for r in results
        for typ, items in (("Preuve (critère couvert)", r["preuves"]),
                           ("Constat (critère manquant)", r.get("constats", [])))
        for p in items
    ], columns=["ID", "Mesure", "Type", "Critère", "Document", "Page", "Extrait"])
    synthese = pd.DataFrame([stats["global"]] + stats["par_theme"])
    nis2 = pd.DataFrame(stats["par_nis2"])

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        synthese.to_excel(writer, sheet_name="Synthèse", index=False)
        matrice_dataframe(results).to_excel(writer, sheet_name="Matrice", index=False)
        preuves.to_excel(writer, sheet_name="Preuves et constats", index=False)
        nis2.to_excel(writer, sheet_name="Vue NIS2", index=False)

        largeurs = {"Synthèse": 18, "Matrice": 40, "Preuves et constats": 45, "Vue NIS2": 18}
        for name, ws in writer.sheets.items():
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = largeurs[name]
                for cell in col[1:]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.freeze_panes = "A2"

        ws = writer.sheets["Matrice"]
        ws.column_dimensions["A"].width = 7
        ws.column_dimensions["D"].width = 22
        for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
            cell = row[0]
            if cell.value in COULEURS:
                fond, texte = COULEURS[cell.value]
                cell.fill = PatternFill("solid", fgColor=fond)
                cell.font = Font(bold=True, color=texte)
    return path


def run_and_export(ids=None, themes=None, progress=None, **eval_kwargs):
    """Enchaîne audit, statistiques, sauvegarde JSON et export Excel."""
    results = run_audit(ids, themes, progress, **eval_kwargs)
    stats = compute_stats(results)
    json_path = save_results(results, stats)
    xlsx_path = export_excel(results, stats, json_path.with_suffix(".xlsx"))
    return results, stats, json_path, xlsx_path
